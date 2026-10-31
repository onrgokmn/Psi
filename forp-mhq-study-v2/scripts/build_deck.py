from pathlib import Path
import json, numpy as np, matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
R=Path(__file__).resolve().parents[1];d=json.loads((R/'results/expanded_summary.json').read_text())
# publication palette: muted navy, teal, sandstone; white body
NAVY=RGBColor(24,49,66); TEAL=RGBColor(27,103,124); INK=RGBColor(40,52,65); MUTED=RGBColor(104,118,129); LIGHT=RGBColor(240,246,247); LINE=RGBColor(222,229,232); ORANGE=RGBColor(174,105,60)
prs=Presentation();prs.slide_width=Inches(13.33);prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]
def shape(sl,x,y,w,h,fill=None,border=None,radius=False):
 sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=fill or RGBColor(255,255,255)
 sh.line.fill.background() if border is None else None
 if border is not None:sh.line.color.rgb=border
 return sh

def txt(sl,text,x,y,w,h,size=20,bold=False,color=INK,align=None,font='Aptos',valign=None):
 t=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=t.text_frame;tf.clear();tf.word_wrap=True;tf.auto_size=MSO_AUTO_SIZE.NONE
 for k,line in enumerate(str(text).split('\n')):
  p=tf.paragraphs[0] if k==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color
  p.space_after=Pt(6);p.space_before=Pt(0)
  if align:p.alignment=align
 return t

def base(section,title,subtitle=''):
 s=prs.slides.add_slide(blank);shape(s,0,0,13.33,.11,NAVY);txt(s,section.upper(),.7,.33,9,.3,10,True,TEAL);txt(s,title,.7,.76,11.9,.83,27,True,NAVY)
 shape(s,.72,1.64,11.87,.015,LINE)
 if subtitle:txt(s,subtitle,.75,1.78,11.9,.7,13,False,MUTED)
 txt(s,'FORP-MHQ   /   ARAŞTIRMA PROTOTİPİ • SENTETİK VERİ',.72,7.1,10,.19,8,False,MUTED)
 txt(s,str(len(prs.slides)).zfill(2),12.2,7.08,.45,.2,9,True,TEAL,PP_ALIGN.RIGHT)
 return s

def bullets(s,items,x=.85,y=2.17,w=11.45,size=19,leading=.83):
 for i,line in enumerate(items):
  shape(s,x,y+i*leading+.12,.065,.065,TEAL)
  txt(s,line,x+.25,y+i*leading,w-.27,leading-.05,size,False,INK)

def card(s,heading,body,x,y,w,h,top=TEAL):
 shape(s,x,y,w,h,LIGHT,radius=True);shape(s,x,y,.06,h,top)
 txt(s,heading,x+.28,y+.19,w-.52,.34,14,True,NAVY);txt(s,body,x+.28,y+.66,w-.58,h-.78,17,False,INK)

def stat(s,value,label,x,y,w=3.0):
 shape(s,x,y,w,1.26,LIGHT,radius=True);txt(s,str(value),x+.18,y+.12,w-.35,.59,30,True,TEAL);txt(s,label,x+.18,y+.82,w-.35,.29,12,False,MUTED)

def image(s,path,x=1.4,y=1.85,w=10.3,h=None):
 s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w) if w else None,height=Inches(h) if h else None)

def twocol(s,lhead,ltxt,rhead,rtxt):
 card(s,lhead,ltxt,.8,2.05,5.7,3.9);card(s,rhead,rtxt,6.83,2.05,5.7,3.9,ORANGE)

# 1
s=prs.slides.add_slide(blank);shape(s,0,0,.17,7.5,TEAL)
txt(s,'UZMANLIK TEZİ  /  GENİŞLETİLMİŞ ARAŞTIRMA SUNUMU',.9,.66,11,.45,13,True,TEAL)
txt(s,'Psikiyatrik bozukluklarda\nnüks korkusu ve\naltı aylık iyi oluş',.88,1.55,11.55,2.9,35,True,NAVY)
txt(s,'Kültürel bağlam · Aile desteği · Kısa ölçüm bataryası',.92,5.12,11.8,.65,19,False,MUTED)
shape(s,.9,6.39,11.5,.017,LINE);txt(s,'Balıkesir Üniversitesi / Ruh Sağlığı ve Hastalıkları AD  •  2026',.9,6.57,11,.45,12,False,INK)
txt(s,'YÖNTEM PROTOTİPİ — GERÇEK HASTA VERİSİ DEĞİLDİR',.9,7.16,11,.22,10,True,ORANGE)
# 2
s=base('Klinik soru','Nüksten korunurken yaşamı kısıtlıyor muyuz?')
txt(s,'Hastanın “yeniden hastalanırsam?” sorusu bazen erken yardım aramayı, bazen sosyal ve mesleki kaçınmayı tetikleyebilir.',.86,2.08,11.3,1.3,23,False,INK)
twocol(s,'Uyum sağlayıcı davranış','Düzenli takip\nErken yardım arama\nTedavi planını sürdürme','İşlev bozucu davranış','Güvence arama\nAşırı belirti izleme\nİş ve ilişkilerden kaçınma')
# 3
s=base('Bilimsel arka plan','Üç yayın, yeni bir araştırma sorusu')
bullets(s,['2023 sistematik derleme: daha düşük yaşam kalitesi; bazı çalışmalarda daha iyi ilaç uyumu.','2025 FORP-MHQ: 865 katılımcıdan geliştirilen 10 maddelik tek faktörlü araç.','2026 transdiagnostik model: korku, utanç, hipervijilans ve kaçınma.','Açık soru: Korku, mevcut belirtilerin ötesinde gelecekteki iyi oluşu açıklar mı?'],size=18,leading=1.07)
# 4
s=base('Literatür boşluğu','Ölçek doğrulaması, prognostik geçerlik değildir')
twocol(s,'Bilinenler','FORP semptom şiddetiyle ve yaşam kalitesiyle ilişkilidir. Kısa bir transdiagnostik ölçüm aracı mevcuttur.','Bilinmeyenler','Altı ay sonraki iyi oluşa artımlı katkı; kültürel anlamlar; destek ve açıklanma endişesinin rolü.')
# 5
s=base('Özgün katkı','Çeviri çalışmasının ötesine geçen tasarım')
bullets(s,['Kısa anketli prospektif klinik kohort','Türkçe FORP-MHQ için izinli, hasta merkezli kültürel uyarlama','Birincil sonlanım: 6. ay WHO-5 (0–100)','Kültürel örüntüler için ayrı kısa sorular ve derin görüşmeler','Takip kaybını hesaba katan duyarlılık analizi'],size=19)
# 6
s=base('Kavramsal model','Korkudan sonraki iyi oluşa giden olası yollar')
card(s,'Nüks korkusu','FORP-MHQ / 10 madde',.8,2.2,3.15,1.55)
card(s,'Düzenleyici bağlam','Aile desteği ve açıklanma endişesi',4.95,2.2,3.45,1.55)
card(s,'İzlem sonucu','Altıncı ay WHO-5',9.1,2.2,3.25,1.55)
shape(s,4.13,2.94,.57,.05,TEAL);shape(s,8.49,2.94,.43,.05,TEAL)
txt(s,'Bağımsız ilişkiyi sınama; nedensellik iddiası yok.',1.14,4.55,11,.8,23,True,NAVY)
# 7
s=base('Birincil amaç','Prospektif, tanılar üstü ilişkiyi ölçmek')
txt(s,'Altıncı ay WHO-5 = f(FORP/10, başlangıç WHO-5, PHQ-4, yaş, cinsiyet, tanı, aile desteği, açıklanma endişesi)',.9,2.4,11.6,1.6,22,True,NAVY)
txt(s,'Ana estimand: 10 FORP puanı başına ayarlanmış WHO-5 ortalama farkı.',.9,4.47,11.3,.75,19,False,INK)
# 8
s=base('Tasarım','Planlanan gerçek klinik araştırma')
for i,(t,detail) in enumerate([('Ön hazırlık','Telif izni + bilişsel görüşme'),('Başlangıç','Kısa ölçek bataryası'),('3. ay','WHO-5'),('6. ay','WHO-5 + klinik sonuç')]):
 x=.8+i*3.13;card(s,t,detail,x,2.4,2.82,2.1)
 if i<3:shape(s,x+2.88,3.42,.18,.04,TEAL)
txt(s,'Etik, kurum izni ve katılımcı onamı olmadan insan araştırması başlatılamaz.',.92,5.66,11.8,.7,15,False,ORANGE)
# 9
s=base('Örneklem','Tek merkezde ardışık hasta alımı')
bullets(s,['18 yaş ve üzeri; geçmiş ruhsal hastalık deneyimi','Depresyon, anksiyete, bipolar, OKB ve uygun diğer tanılar','Klinik karar verme yetisini/yanıtı bozacak akut durumların kayıtlı değerlendirmesi','Uygunluk, ret ve takip kaybını içeren denetlenebilir hasta akışı','Tanılar arası karşılaştırma öncesi ölçüm değişmezliği kontrolü'],size=17,leading=.78)
#10
s=base('Hasta yükü','İlk görüşmede 25–30 kısa madde')
for i,(v,l) in enumerate([('10','FORP-MHQ'),('4','PHQ-4'),('5','WHO-5'),('4','Aile desteği'),('2','Kültürel sorular'),('5*','İlaç uyumu')]):
 stat(s,v,l,.77+(i%3)*4.1,2.12+(i//3)*1.65,3.63)
txt(s,'*Yalnızca ilaç kullananlarda. Hedef süre 7–10 dk; henüz klinikte ölçülmedi.',.84,5.96,11.8,.52,13,False,MUTED)
#11
s=base('Kültür','Stereotip değil, hastanın bildirdiği deneyim')
twocol(s,'Nicel göstergeler','MSPSS aile desteği\nAçıklanma endişesi\nEn çok korkulan sonuç','Nitel sorgulama','Aile beklentileri\nMahremiyet ve açıklama\nÖzerklik, iş, ilişki ve anlam')
#12
s=base('Ölçek','FORP-MHQ: izinli kültürel uyarlama adımları')
bullets(s,['İki ileri çeviri ve uzlaşma','Kör geri çeviri, uzman paneli','Özellikle “nüks” ve “kötüleşme” kavramlarında bilişsel görüşme','Pilot zamanlama ve madde anlaşılırlığı','Ordinal DFA/WLSMV, güvenirlik, hipotez temelli geçerlik'],size=18,leading=.84)
#13
s=base('Psikometri','Alfa tek başına geçerlik kanıtı değildir')
card(s,'İçerik','Hastalar maddeleri ne anlıyor?',.78,2.13,3.77,2.26)
card(s,'Yapı','Tek faktör, eşikler ve yükler',4.77,2.13,3.77,2.26)
card(s,'Tekrar ölçüm','ICC, klinik stabilite',8.76,2.13,3.77,2.26)
txt(s,'Simüle madde kodlarından Türkçe validasyon sonucu çıkarılamaz.',.84,5.21,11.55,.8,19,True,ORANGE)
#14
s=base('Nitel kol','Hastanın anlam dünyasını yakalamak')
bullets(s,['Yaklaşık 15–20 gönüllü, amaçlı çeşitlilik örneklemesi','Nüks korkusu, aile, mahremiyet ve kaçınma üzerine yarı yapılandırılmış görüşmeler','Transkript, kodlama, refleksif tematik analiz ve COREQ raporu','Nitel veri bulunmadığından temalar ve hasta alıntıları üretilmedi'],size=18,leading=1.05)
#15
s=base('Veri kaynağı','Hesaplar sabit tohumlu sentetik kohortta')
stat(s,'360','Oluşturulan kayıt',.8,2.1,3.7);stat(s,'303','6. ay WHO-5',4.84,2.1,3.7);stat(s,'65','Retest çifti',8.88,2.1,3.7)
txt(s,'Modelin FORP ile altıncı ay WHO-5 arasındaki negatif ilişkiyi açıkça üretmesi, p değerinin klinik kanıt olarak kullanılmasını engeller.',.95,4.26,11.5,1.2,18,False,INK)
#16
s=base('Hasta akışı','Altı aylık veri tamamlama oranı')
image(s,R/'figures/cohort_flow.png',2.65,1.88,7.8)
#17
s=base('Başlangıç','Temel demografik/klinik özellikler')
stat(s,'42.1','Yaş ortalaması',.8,2.17,3.75);stat(s,'61.7%','Kadın kayıt oranı',4.78,2.17,3.75);stat(s,'18.28','FORP ortalaması',8.76,2.17,3.75)
txt(s,'Tanı dağılımı: depresyon 122 • anksiyete 121 • bipolar 47 • OKB 46 • diğer 24',.85,4.32,11.7,.9,18)
#18
s=base('Dağılım','FORP puanlarının sentetik dağılımı')
image(s,R/'figures/forp_distribution.png',1.25,1.89,10.5)
#19
s=base('İzlem','Takip kaybı yüksek korkuyla ilişkili üretildi')
stat(s,'57','6. ay kayıp kayıt',.85,2.15,3.65);stat(s,'+3.69','Kayıplarda FORP farkı',4.85,2.15,3.65);stat(s,'−5.12','Kayıplarda WHO-5 farkı',8.85,2.15,3.65)
txt(s,'Bu özellik veri üretim algoritmasında yer alır. Klinik bir gözlem değildir.',.91,4.43,11.4,.85,19,False,ORANGE)
#20
s=base('İzlem','Tanı gruplarına göre gözlenen sonuç oranları')
image(s,R/'figures/attrition_by_diagnosis.png',1.2,1.9,10.75)
#21
s=base('Güvenirlik','Simülasyonun psikometrik sayıları')
stat(s,'0.948','Cronbach alfa',.8,2.22,3.7);stat(s,'0.948','Sürekli ML omega (yaklaşık)',4.83,2.22,3.7);stat(s,'0.954','ICC (n=65)',8.86,2.22,3.7)
txt(s,'Yüksek değerler üretilmiş ortak faktörün sonucudur. Yetkili Türkçe form için gerçek faktör analizi henüz yok.',.89,4.43,11.5,1.02,18,False,INK)
#22
s=base('Birincil sonuç','Düzeltilmiş ilişki: −2.13 WHO-5 puanı')
txt(s,'−2.13',1.05,2.04,7.3,1.63,60,True,TEAL)
txt(s,'FORP toplamında 10 puan artış başına',1.13,3.79,9.3,.65,22,False,INK)
txt(s,'%95 GA: −3.97 ile −0.30     •     p=0.023',1.13,4.81,10.3,.66,20,True,NAVY)
txt(s,'Hesaplanan etki simülasyonda tasarım gereği vardır.',1.13,6.05,11,.4,14,False,ORANGE)
#23
s=base('Ana görsel','Düzeltilmiş ilişki: bütünlük ve belirsizlik')
image(s,R/'figures/prospective_association.png',1.0,1.86,11.15)
#24
s=base('Artımlı açıklayıcılık','İstatistiksel anlamlılık ≠ klinik öngörü')
stat(s,'0.470','Temel model R²',.83,2.22,3.69);stat(s,'0.482','FORP eklenen R²',4.83,2.22,3.69);stat(s,'0.011','ΔR²',8.83,2.22,3.69)
txt(s,'Ek açıklanan varyans yaklaşık 1,1 yüzde puan. Klinik değer, bağımsız doğrulama ve karar analizi gerektirir.',.92,4.49,11.5,1.3,20,False,INK)
#25
s=base('Duyarlılık','Takip kaybı için ters olasılık ağırlığı')
txt(s,'Ağırlıklandırılmış katsayı: −2.13',1.05,2.22,10.5,.73,29,True,TEAL)
txt(s,'%95 GA −3.97 ile −0.29    •    p=0.023',1.08,3.22,10.6,.7,21,False,INK)
txt(s,'Kayıp mekanizması modelden bilindiğinde beklenen benzerlik: gerçek klinikte varsayımlar belirsizdir.',1.07,4.54,10.8,1.1,18,False,MUTED)
#26
s=base('MNAR senaryoları','Kayıp kayıtların olası sonuçlarıyla sınama')
vals=d['delta_sensitivity']; X=[z['delta'] for z in vals];Y=[z['beta'] for z in vals]
fig,ax=plt.subplots(figsize=(10.5,4.4),dpi=170);ax.plot(X,Y,marker='o',color='#1b677c',linewidth=2.5);ax.axhline(0,color='#c8d7dc',linestyle='--');ax.set_xlabel('Missing-outcome delta shift (WHO-5 points)');ax.set_ylabel('Adjusted FORP coefficient');ax.set_title('Assumption-driven pattern-mixture check',loc='left');ax.spines[['top','right']].set_visible(False);fig.tight_layout();fig.savefig(R/'figures/delta_sensitivity.png');plt.close(fig)
image(s,R/'figures/delta_sensitivity.png',1.4,1.88,10.25)
#27
s=base('Kültürel etkileşim','Aile desteği moderasyonu: desteklenmedi')
stat(s,'−0.06','Etkileşim katsayısı',.87,2.22,5.5);stat(s,'0.937','p değeri',6.9,2.22,5.5)
txt(s,'Bu, Türk hastalarda moderasyon olmadığı anlamına gelmez: veri üretim sürecinde gerçek kültürel deneyim yoktur.',.91,4.47,11.6,1.4,19,False,INK)
#28
s=base('Klinik olay','Oluşturulmuş kötüleşme ≠ klinik nüks')
stat(s,'51 / 303','Yapay kötüleşme bayrağı',.86,2.2,5.5);stat(s,'2.35','FORP/10 için OR',6.91,2.2,5.5)
txt(s,'Epizod tanısı, başlangıç remisyonu, olay tarihi ve bağımsız klinisyen doğrulaması olmadan “nüks riski” sonucu çıkmaz.',.9,4.51,11.6,1.2,19,False,INK)
#29
s=base('Yorum','Bu araştırmada henüz neyi bilmiyoruz?')
bullets(s,['Türkçe maddelerin hastalarca nasıl anlaşıldığını','Gerçek FORP puanı dağılımını ve faktör yapısını','Aile etkisinin yönünü ve büyüklüğünü','Altı aylık gerçek klinik kötüleşme oranını','Klinik olarak anlamlı öngörü farkını'],size=18,leading=.83)
#30
s=base('Sınırlılıklar','Modeli klinik kanıt gibi okumamak gerekir')
bullets(s,['Verinin tamamı yazılımla üretildi','FORP–WHO-5 ilişkisi denklemde kuruldu','Klinik tanılar ve olaylar doğrulanmış değil','Türkçe FORP-MHQ izni/uyarlaması henüz yok','DFA ve görüşme temaları mevcut değil'],size=19,leading=.82)
#31
s=base('Güçlü tasarım','Gerçek çalışmayı iyi yapan beş karar')
bullets(s,['İlk olarak içerik geçerliği ve hasta bilişsel görüşmeleri','Önceden kaydedilmiş tek bir birincil estimand','Tanı özelinde klinik sonlanım ve kör sonuç değerlendirmesi','Tek merkezli ardışık örnekleme ve takip yönetimi','Eksik veride açık varsayımlar ve duyarlılık analizleri'],size=18,leading=.86)
#32
s=base('İş akışı','Klinik uygulama öncesi yapılacaklar')
for i,(head,body) in enumerate([('01','Telif izni'),('02','Etik/kurum'),('03','Pilot'),('04','Güç hesabı'),('05','Takip + analiz')]):
 x=.82+i*2.51;shape(s,x,2.5,2.28,2.45,LIGHT,radius=True);txt(s,head,x+.2,2.73,1.8,.59,26,True,TEAL);txt(s,body,x+.2,3.72,1.9,.6,18,True,NAVY)
#33
s=base('Yayın stratejisi','Q1 düzeyinde gerçek katkı için gerekli eşikler')
bullets(s,['Özgün ve klinik açıdan önemli bir birincil soru','COSMIN 2.0 + STROBE raporlaması','Yeterli güç ve gerçek takip; örneklem akışının şeffaflığı','Artımlı katkı ve klinik anlamlılığın birlikte değerlendirilmesi','Gerçek nitel verilerle kültürel bütünleştirme'],size=18,leading=.85)
#34
s=base('Kaynaklar','Kuramsal ve metodolojik temel')
bullets(s,['Coutts-Bain et al. Clin Psychol Rev. 2023. doi:10.1016/j.cpr.2023.102342','Coutts-Bain et al. Br J Clin Psychol. 2025. doi:10.1111/bjc.12536','Coutts-Bain et al. J Psychopathol Clin Sci. 2026. doi:10.1037/abn0001040','Gagnier et al. Qual Life Res. 2025. doi:10.1007/s11136-025-03950-x','von Elm et al. PLoS Med. 2007. doi:10.1371/journal.pmed.0040296'],size=16,leading=.84)
#35
s=base('Özet','Nüksten korkmak ile yaşamı ertelemek arasındaki fark')
txt(s,'Klinik soru değişmiyor:',1.05,2.27,10.9,.61,18,True,TEAL)
txt(s,'Korku, yalnızca bugünkü belirtilerin yansıması mı; yoksa sonraki iyilik haline ayrıca ışık tutuyor mu?',1.04,3.08,11.12,2.1,26,True,NAVY)
txt(s,'Gerçek hasta verisi olmadan yanıt veremeyiz. Bu paket, bu soruyu doğru biçimde araştırmak için kurulmuş bir yöntem altyapısıdır.',1.06,5.68,11.4,.9,16,False,MUTED)
# notes
prs.core_properties.title='Psychiatric Relapse Fear - Thesis Defense (Methodological Demonstration)'
prs.core_properties.subject='Simulated transdiagnostic research study'
prs.core_properties.author='Research working draft'
path=R/'slides/defense_full.pptx';prs.save(path)
print('Slides created:',len(prs.slides),path)